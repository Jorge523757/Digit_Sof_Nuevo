"""
DIGIT SOFT - Generador de Reportes para Órdenes de Servicio
Clase para generar reportes en Excel y PDF
"""

from io import BytesIO
from datetime import datetime
from django.db.models import Count, Sum, Q, Avg
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.pdfgen import canvas
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

from .models import OrdenServicio


class GeneradorReportesOrdenes:
    """Generador profesional de reportes de órdenes de servicio"""

    def __init__(self):
        if REPORTLAB_AVAILABLE:
            self.estilos = getSampleStyleSheet()
            self.estilos.add(ParagraphStyle(
                name='CustomTitle',
                parent=self.estilos['Heading1'],
                fontSize=18,
                textColor=colors.HexColor('#1e3c72'),
                spaceAfter=30,
                alignment=1  # Centrado
            ))

    def filtrar_ordenes(self, filtros):
        """
        Filtrar órdenes según los criterios proporcionados
        """
        queryset = OrdenServicio.objects.select_related(
            'cliente',
            'tecnico_asignado'
        ).all()

        # Filtro por rango de fechas
        if filtros.get('fecha_desde'):
            queryset = queryset.filter(fecha_recepcion__gte=filtros['fecha_desde'])

        if filtros.get('fecha_hasta'):
            queryset = queryset.filter(fecha_recepcion__lte=filtros['fecha_hasta'])

        # Filtro por estado
        if filtros.get('estado'):
            queryset = queryset.filter(estado=filtros['estado'])

        # Filtro por prioridad
        if filtros.get('prioridad'):
            queryset = queryset.filter(prioridad=filtros['prioridad'])

        # Filtro por cliente
        if filtros.get('cliente_id'):
            queryset = queryset.filter(cliente_id=filtros['cliente_id'])

        # Filtro por técnico
        if filtros.get('tecnico_id'):
            queryset = queryset.filter(tecnico_asignado_id=filtros['tecnico_id'])

        # Filtro por tipo de equipo
        if filtros.get('tipo_equipo'):
            queryset = queryset.filter(tipo_equipo__icontains=filtros['tipo_equipo'])

        # Ordenar
        queryset = queryset.order_by('-fecha_recepcion')

        return queryset

    def generar_excel(self, filtros, incluir_estadisticas=True):
        """
        Generar reporte en Excel profesional
        """
        ordenes = self.filtrar_ordenes(filtros)

        # Crear workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Órdenes de Servicio"

        # Estilos
        header_fill = PatternFill(start_color="1e3c72", end_color="1e3c72", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=12)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Título
        ws.merge_cells('A1:K1')
        titulo_cell = ws['A1']
        titulo_cell.value = "REPORTE DE ÓRDENES DE SERVICIO - DIGIT SOFT"
        titulo_cell.font = Font(size=16, bold=True, color="1e3c72")
        titulo_cell.alignment = Alignment(horizontal='center', vertical='center')

        # Fecha de generación
        ws.merge_cells('A2:K2')
        fecha_cell = ws['A2']
        fecha_cell.value = f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        fecha_cell.alignment = Alignment(horizontal='center')
        fecha_cell.font = Font(italic=True)

        # Filtros aplicados
        fila_actual = 4
        ws.merge_cells(f'A{fila_actual}:K{fila_actual}')
        filtros_cell = ws[f'A{fila_actual}']
        filtros_cell.value = "FILTROS APLICADOS:"
        filtros_cell.font = Font(bold=True, size=11)

        fila_actual += 1
        filtros_texto = []
        if filtros.get('fecha_desde'):
            filtros_texto.append(f"Desde: {filtros['fecha_desde']}")
        if filtros.get('fecha_hasta'):
            filtros_texto.append(f"Hasta: {filtros['fecha_hasta']}")
        if filtros.get('estado'):
            filtros_texto.append(f"Estado: {filtros['estado']}")
        if filtros.get('prioridad'):
            filtros_texto.append(f"Prioridad: {filtros['prioridad']}")

        if filtros_texto:
            ws.merge_cells(f'A{fila_actual}:K{fila_actual}')
            ws[f'A{fila_actual}'].value = " | ".join(filtros_texto)
            ws[f'A{fila_actual}'].font = Font(italic=True, size=10)
        else:
            ws.merge_cells(f'A{fila_actual}:K{fila_actual}')
            ws[f'A{fila_actual}'].value = "Todas las órdenes"
            ws[f'A{fila_actual}'].font = Font(italic=True, size=10)

        # Encabezados
        fila_actual += 2
        headers = ['N° Orden', 'Fecha', 'Cliente', 'Técnico', 'Equipo', 'Marca', 'Modelo',
                   'Prioridad', 'Estado', 'Costo Total', 'Observaciones']

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=fila_actual, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border

        # Datos
        fila_actual += 1
        for orden in ordenes:
            cliente_nombre = f"{orden.cliente.nombres} {orden.cliente.apellidos}" if orden.cliente else "N/A"
            tecnico_nombre = f"{orden.tecnico_asignado.nombres} {orden.tecnico_asignado.apellidos}" if orden.tecnico_asignado else "Sin asignar"

            row_data = [
                orden.numero_orden,
                orden.fecha_recepcion.strftime('%d/%m/%Y') if orden.fecha_recepcion else "",
                cliente_nombre,
                tecnico_nombre,
                orden.tipo_equipo or "",
                orden.marca or "",
                orden.modelo or "",
                orden.get_prioridad_display() if orden.prioridad else "",
                orden.get_estado_display() if orden.estado else "",
                float(orden.costo_total) if orden.costo_total else 0,
                orden.observaciones[:50] if orden.observaciones else ""
            ]

            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=fila_actual, column=col_num)
                cell.value = value
                cell.border = border

                if col_num == 10:  # Costo Total
                    cell.number_format = '$#,##0.00'
                    cell.alignment = Alignment(horizontal='right')
                else:
                    cell.alignment = Alignment(horizontal='left', vertical='center')

            fila_actual += 1

        # Estadísticas
        if incluir_estadisticas:
            fila_actual += 2
            ws.merge_cells(f'A{fila_actual}:K{fila_actual}')
            stats_cell = ws[f'A{fila_actual}']
            stats_cell.value = "ESTADÍSTICAS"
            stats_cell.font = Font(bold=True, size=12, color="1e3c72")
            stats_cell.alignment = Alignment(horizontal='center')

            fila_actual += 1

            # Calcular estadísticas
            total_ordenes = ordenes.count()
            total_costo = ordenes.aggregate(Sum('costo_total'))['costo_total__sum'] or 0
            promedio_costo = ordenes.aggregate(Avg('costo_total'))['costo_total__avg'] or 0

            estados_stats = ordenes.values('estado').annotate(count=Count('id')).order_by('-count')
            prioridades_stats = ordenes.values('prioridad').annotate(count=Count('id')).order_by('-count')

            stats = [
                ['Total de Órdenes:', total_ordenes],
                ['Costo Total:', f'${total_costo:,.2f}'],
                ['Promedio por Orden:', f'${promedio_costo:,.2f}'],
            ]

            for stat in stats:
                ws.cell(row=fila_actual, column=1).value = stat[0]
                ws.cell(row=fila_actual, column=1).font = Font(bold=True)
                ws.cell(row=fila_actual, column=2).value = stat[1]
                fila_actual += 1

        # Ajustar anchos de columna
        column_widths = [15, 12, 25, 25, 15, 15, 15, 12, 18, 15, 30]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        # Guardar en memoria
        excel_file = BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)

        return excel_file

    def generar_pdf(self, filtros, orientacion='portrait'):
        """
        Generar reporte en PDF profesional
        """
        if not REPORTLAB_AVAILABLE:
            raise ImportError("ReportLab no está instalado. Ejecuta: pip install reportlab")

        ordenes = self.filtrar_ordenes(filtros)

        # Crear PDF en memoria
        pdf_buffer = BytesIO()

        # Configurar página
        pagesize = A4 if orientacion == 'portrait' else landscape(A4)
        doc = SimpleDocTemplate(pdf_buffer, pagesize=pagesize,
                                rightMargin=30, leftMargin=30,
                                topMargin=30, bottomMargin=30)

        # Contenido
        elementos = []

        # Título
        titulo = Paragraph("REPORTE DE ÓRDENES DE SERVICIO", self.estilos['CustomTitle'])
        elementos.append(titulo)
        elementos.append(Spacer(1, 12))

        # Subtítulo
        fecha_gen = datetime.now().strftime('%d/%m/%Y %H:%M')
        subtitulo = Paragraph(f"Generado el: {fecha_gen}", self.estilos['Normal'])
        elementos.append(subtitulo)
        elementos.append(Spacer(1, 20))

        # Filtros
        filtros_texto = "Filtros aplicados: "
        filtros_list = []
        if filtros.get('fecha_desde'):
            filtros_list.append(f"Desde {filtros['fecha_desde']}")
        if filtros.get('fecha_hasta'):
            filtros_list.append(f"Hasta {filtros['fecha_hasta']}")
        if filtros.get('estado'):
            filtros_list.append(f"Estado: {filtros['estado']}")
        if filtros.get('prioridad'):
            filtros_list.append(f"Prioridad: {filtros['prioridad']}")

        if filtros_list:
            filtros_texto += " | ".join(filtros_list)
        else:
            filtros_texto += "Todas las órdenes"

        filtros_p = Paragraph(filtros_texto, self.estilos['Italic'])
        elementos.append(filtros_p)
        elementos.append(Spacer(1, 20))

        # Tabla de datos
        datos_tabla = [
            ['N° Orden', 'Fecha', 'Cliente', 'Técnico', 'Equipo', 'Estado', 'Costo']
        ]

        for orden in ordenes[:100]:  # Limitar a 100 para el PDF
            cliente_nombre = f"{orden.cliente.nombres[:15]}" if orden.cliente else "N/A"
            tecnico_nombre = f"{orden.tecnico_asignado.nombres[:15]}" if orden.tecnico_asignado else "N/A"
            equipo_completo = f"{orden.tipo_equipo} {orden.marca}".strip()[:20] if orden.tipo_equipo else "N/A"

            datos_tabla.append([
                orden.numero_orden[:12],
                orden.fecha_recepcion.strftime('%d/%m/%Y') if orden.fecha_recepcion else 'N/A',
                cliente_nombre,
                tecnico_nombre,
                equipo_completo or 'N/A',
                orden.get_estado_display()[:20] if orden.estado else 'N/A',
                f"${orden.costo_total:,.0f}" if orden.costo_total else '$0'
            ])

        tabla = Table(datos_tabla)
        tabla.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3c72')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey])
        ]))

        elementos.append(tabla)

        # Estadísticas
        elementos.append(Spacer(1, 20))
        stats_title = Paragraph("<b>ESTADÍSTICAS</b>", self.estilos['Heading2'])
        elementos.append(stats_title)
        elementos.append(Spacer(1, 10))

        total_ordenes = ordenes.count()
        total_costo = ordenes.aggregate(Sum('costo_total'))['costo_total__sum'] or 0

        stats_text = f"""
        <b>Total de órdenes:</b> {total_ordenes}<br/>
        <b>Costo total:</b> ${total_costo:,.2f}<br/>
        <b>Promedio por orden:</b> ${total_costo/total_ordenes if total_ordenes > 0 else 0:,.2f}
        """

        stats_p = Paragraph(stats_text, self.estilos['Normal'])
        elementos.append(stats_p)

        # Generar PDF
        doc.build(elementos)

        pdf_buffer.seek(0)
        return pdf_buffer

