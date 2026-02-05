"""
DIGT SOFT - Servicios para Generación de Reportes
"""

from io import BytesIO
from datetime import datetime
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from django.conf import settings
import os


class GeneradorReportePDF:
    """Genera reportes en formato PDF"""
    
    @staticmethod
    def generar_reporte_dano(registro):
        """Genera PDF del reporte de daño"""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        elementos = []
        
        # Estilos
        estilos = getSampleStyleSheet()
        estilo_titulo = ParagraphStyle(
            'CustomTitle',
            parent=estilos['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#1e3c72'),
            spaceAfter=30,
            alignment=TA_CENTER
        )
        
        estilo_subtitulo = ParagraphStyle(
            'CustomSubtitle',
            parent=estilos['Heading2'],
            fontSize=16,
            textColor=colors.HexColor('#2a5298'),
            spaceAfter=12
        )
        
        estilo_normal = ParagraphStyle(
            'CustomNormal',
            parent=estilos['Normal'],
            fontSize=11,
            spaceAfter=12
        )
        
        # Título
        elementos.append(Paragraph("DIGIT SOFT", estilo_titulo))
        elementos.append(Paragraph("Reporte de Daño de Equipo", estilo_subtitulo))
        elementos.append(Spacer(1, 0.3*inch))
        
        # Información del reporte
        datos_tabla = [
            ['Número de Factura:', registro.numero_factura or 'Sin asignar'],
            ['Fecha de Reporte:', registro.fecha_reporte.strftime('%d/%m/%Y %H:%M')],
            ['Usuario:', str(registro.usuario.get_full_name() or registro.usuario.username)],
        ]
        
        tabla_info = Table(datos_tabla, colWidths=[2.5*inch, 4*inch])
        tabla_info.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#e8f4f8')),
            ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#1e3c72')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
        ]))
        
        elementos.append(tabla_info)
        elementos.append(Spacer(1, 0.3*inch))
        
        # Descripción del daño
        elementos.append(Paragraph("Descripción del Daño:", estilo_subtitulo))
        elementos.append(Paragraph(registro.descripcion_dano, estilo_normal))
        elementos.append(Spacer(1, 0.3*inch))
        
        # Evidencia fotográfica
        imagenes = registro.imagenes.all()
        if imagenes:
            elementos.append(Paragraph(f"Evidencia Fotográfica ({imagenes.count()} imagen{'es' if imagenes.count() > 1 else ''}):", estilo_subtitulo))
            for img in imagenes:
                try:
                    ruta_imagen = img.imagen.path
                    if os.path.exists(ruta_imagen):
                        imagen_rl = Image(ruta_imagen, width=3*inch, height=2*inch)
                        elementos.append(imagen_rl)
                        if img.es_principal:
                            elementos.append(Paragraph("(Imagen Principal)", estilo_normal))
                        elementos.append(Spacer(1, 0.2*inch))
                except Exception as e:
                    elementos.append(Paragraph(f"Error al cargar imagen: {str(e)}", estilo_normal))
        
        # Pie de página
        elementos.append(Spacer(1, 0.5*inch))
        elementos.append(Paragraph(
            f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            ParagraphStyle('Footer', parent=estilos['Normal'], fontSize=8, textColor=colors.grey, alignment=TA_CENTER)
        ))
        
        # Construir PDF
        doc.build(elementos)
        buffer.seek(0)
        return buffer


class GeneradorReporteExcel:
    """Genera reportes en formato Excel"""
    
    @staticmethod
    def generar_reporte_dano(registro):
        """Genera Excel del reporte de daño"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Reporte de Daño"
        
        # Estilos
        titulo_font = Font(name='Arial', size=16, bold=True, color='1E3C72')
        subtitulo_font = Font(name='Arial', size=12, bold=True, color='2A5298')
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        normal_font = Font(name='Arial', size=10)
        
        header_fill = PatternFill(start_color='1E3C72', end_color='1E3C72', fill_type='solid')
        data_fill = PatternFill(start_color='E8F4F8', end_color='E8F4F8', fill_type='solid')
        
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Título
        ws.merge_cells('A1:D1')
        ws['A1'] = 'DIGIT SOFT - Reporte de Daño de Equipo'
        ws['A1'].font = titulo_font
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 30
        
        # Información del reporte
        fila = 3
        
        # Encabezados
        ws[f'A{fila}'] = 'Campo'
        ws[f'A{fila}'].font = header_font
        ws[f'A{fila}'].fill = header_fill
        ws[f'A{fila}'].alignment = Alignment(horizontal='center')
        ws[f'A{fila}'].border = border
        
        ws[f'B{fila}'] = 'Valor'
        ws[f'B{fila}'].font = header_font
        ws[f'B{fila}'].fill = header_fill
        ws[f'B{fila}'].alignment = Alignment(horizontal='center')
        ws[f'B{fila}'].border = border
        
        # Datos
        datos = [
            ('Número de Factura', registro.numero_factura or 'Sin asignar'),
            ('Fecha de Reporte', registro.fecha_reporte.strftime('%d/%m/%Y %H:%M')),
            ('Usuario', str(registro.usuario.get_full_name() or registro.usuario.username)),
            ('Descripción del Daño', registro.descripcion_dano),
        ]
        
        imagenes = registro.imagenes.all()
        if imagenes:
            datos.append(('Evidencias Fotográficas', f'{imagenes.count()} imagen{"es" if imagenes.count() > 1 else ""}'))
        
        fila += 1
        for campo, valor in datos:
            ws[f'A{fila}'] = campo
            ws[f'A{fila}'].font = Font(name='Arial', size=10, bold=True)
            ws[f'A{fila}'].fill = data_fill
            ws[f'A{fila}'].border = border
            
            ws[f'B{fila}'] = valor
            ws[f'B{fila}'].font = normal_font
            ws[f'B{fila}'].border = border
            ws[f'B{fila}'].alignment = Alignment(wrap_text=True, vertical='top')
            
            # Ajustar altura si es descripción
            if campo == 'Descripción del Daño':
                ws.row_dimensions[fila].height = max(30, len(valor) / 2)
            
            fila += 1
        
        # Ajustar anchos de columna
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 60
        
        # Pie de página
        fila += 2
        ws[f'A{fila}'] = f'Generado el {datetime.now().strftime("%d/%m/%Y %H:%M")}'
        ws[f'A{fila}'].font = Font(name='Arial', size=8, color='808080')
        
        # Guardar en buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
    
    @staticmethod
    def generar_lista_reportes(reportes):
        """Genera Excel con lista de reportes"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Lista de Reportes"
        
        # Estilos
        titulo_font = Font(name='Arial', size=14, bold=True, color='1E3C72')
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1E3C72', end_color='1E3C72', fill_type='solid')
        
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Título
        ws.merge_cells('A1:F1')
        ws['A1'] = f'Lista de Reportes de Daños - {datetime.now().strftime("%d/%m/%Y")}'
        ws['A1'].font = titulo_font
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 25
        
        # Encabezados
        headers = ['No. Factura', 'Usuario', 'Fecha', 'Descripción', 'Imágenes', 'Estado']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        
        # Datos
        fila = 4
        for reporte in reportes:
            ws.cell(row=fila, column=1, value=reporte.numero_factura or 'Sin asignar')
            ws.cell(row=fila, column=2, value=str(reporte.usuario.get_full_name() or reporte.usuario.username))
            ws.cell(row=fila, column=3, value=reporte.fecha_reporte.strftime('%d/%m/%Y %H:%M'))
            ws.cell(row=fila, column=4, value=reporte.descripcion_dano[:100] + '...' if len(reporte.descripcion_dano) > 100 else reporte.descripcion_dano)
            ws.cell(row=fila, column=5, value=reporte.imagenes.count())
            ws.cell(row=fila, column=6, value='Activo')
            
            # Aplicar bordes
            for col in range(1, 7):
                ws.cell(row=fila, column=col).border = border
                ws.cell(row=fila, column=col).alignment = Alignment(wrap_text=True, vertical='top')
            
            fila += 1
        
        # Ajustar anchos
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 50
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 12
        
        # Guardar
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

